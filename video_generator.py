import io
import os
import json
import re
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
import requests
import streamlit as st
from pypdf import PdfReader
from openpyxl import load_workbook
from openai import OpenAI
from runwayml import RunwayML, TaskFailedError

APP_TITLE = "Agentic Foresight Film Lab"
DEFAULT_MODEL = "gpt-5.4"
DEFAULT_VIDEO_MODEL = "gen4.5"
RATIO = "1280:720"

ANALYST_PROMPT = """
You are the FORECAST ANALYST in a human-in-the-loop strategic foresight system.
Read the supplied forecast faithfully. Do not invent trends, evidence, demographics,
claims or implications that are not supported by the source.

Return JSON only with:
{
  "title": "...",
  "central_shift": "...",
  "drivers": ["..."],
  "tensions": ["..."],
  "consumer_implications": ["..."],
  "strategic_implications": ["..."],
  "evidence_notes": ["..."],
  "visual_symbols": ["..."],
  "uncertainties": ["..."]
}

Keep evidence_notes traceable to the source. Put weak/ambiguous claims in uncertainties.
"""

CREATIVE_PROMPT = """
You are an experimental fashion-film creative director translating strategic foresight
into an ABSTRACT VISUAL FILM.

Avoid: PowerPoint aesthetics, talking heads, voiceover, literal infographics, generic
corporate stock imagery, logos, long captions, and merely illustrating sentences.

Prefer: visual metaphor, materiality, macro texture, surreal transformations, bodies
without identifiable faces, objects, architecture, speculative retail spaces, analogue
imperfection, digital artefacts, fashion-editorial composition, light, shadow, motion,
rhythm, tension and contrast.

Create a coherent film language, not disconnected AI clips.

Return JSON only:
{
  "concept_title": "...",
  "creative_thesis": "...",
  "visual_world": "...",
  "palette_language": "...",
  "material_language": ["..."],
  "camera_language": ["..."],
  "motion_language": ["..."],
  "sound_direction": "...",
  "typography_rule": "...",
  "avoid": ["..."]
}
"""

STORYBOARD_PROMPT = """
You are the STORYBOARD AGENT. Create exactly {scene_count} scenes for a silent,
abstract strategic-foresight film. The film must communicate the approved forecast
through visual metaphor while remaining faithful to it.

Return JSON only:
{{
  "scenes": [
    {{
      "scene": 1,
      "purpose": "opening|driver|tension|shift|implication|closing",
      "forecast_anchor": "short source-grounded idea this scene communicates",
      "visual_metaphor": "what the idea becomes visually",
      "shot": "precise cinematic shot description",
      "motion": "precise movement/transformation",
      "on_screen_text": "0-5 words, or empty string",
      "duration_sec": 5,
      "runway_prompt": "complete standalone video-generation prompt"
    }}
  ]
}}

Rules:
- No voiceover.
- On-screen text is optional and must be 0-5 words.
- Each Runway prompt must be visual, cinematic and self-contained.
- Maintain continuity in materials, lighting, camera language and world.
- Do not use named living artists/directors as style references.
- Do not introduce claims absent from the approved forecast.
"""

CRITIC_PROMPT = """
You are the CRITIC AGENT in a human-in-the-loop foresight-film system.
Compare the storyboard with the approved forecast and creative direction.

Return JSON only:
{
  "faithfulness_score": 0,
  "visual_coherence_score": 0,
  "abstraction_score": 0,
  "issues": ["..."],
  "recommended_changes": ["..."]
}

Scores are 0-100. Penalise invented claims, generic stock-video language, repetition,
literal infographics, long text, and visual ideas that do not communicate the forecast.
"""


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((p.extract_text() or "").strip() for p in reader.pages if (p.extract_text() or "").strip())


def extract_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    chunks = []
    for ws in wb.worksheets:
        chunks.append(f"\nSHEET: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                chunks.append(" | ".join(vals))
    return "\n".join(chunks)


def clean_text(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_json(text: str):
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return json.loads(text)


def llm_json(client, model, system_prompt, payload):
    response = client.responses.create(
        model=model,
        instructions=system_prompt,
        input=json.dumps(payload, ensure_ascii=False),
    )
    return parse_json(response.output_text)


def run_agents(client, model, forecast_text, scene_count):
    analysis = llm_json(
        client, model, ANALYST_PROMPT,
        {"forecast_document": forecast_text}
    )
    creative = llm_json(
        client, model, CREATIVE_PROMPT,
        {"approved_forecast_analysis": analysis}
    )
    storyboard = llm_json(
        client, model, STORYBOARD_PROMPT.format(scene_count=scene_count),
        {"approved_forecast_analysis": analysis, "creative_direction": creative}
    )
    critique = llm_json(
        client, model, CRITIC_PROMPT,
        {
            "approved_forecast_analysis": analysis,
            "creative_direction": creative,
            "storyboard": storyboard
        }
    )
    return analysis, creative, storyboard, critique


def generate_runway_clip(runway_client, prompt, duration):
    # Gen-4.5 currently supports short generations. Clamp the UI value safely.
    duration = int(max(2, min(10, duration)))
    task = runway_client.image_to_video.create(
        model=DEFAULT_VIDEO_MODEL,
        prompt_text=prompt,
        ratio=RATIO,
        duration=duration,
    ).wait_for_task_output()
    if not task.output:
        raise RuntimeError("Runway returned no video URL.")
    return task.output[0]


def download(url, destination):
    r = requests.get(url, timeout=120)
    r.raise_for_status()
    Path(destination).write_bytes(r.content)


def storyboard_df(storyboard):
    return pd.DataFrame(storyboard["scenes"])


st.set_page_config(page_title=APP_TITLE, layout="wide")
st.title(APP_TITLE)
st.caption("Forecast → analyst agent → creative director → storyboard → critic → human approval → generative video")

with st.sidebar:
    st.header("API")
    openai_key = st.text_input("OpenAI API key", type="password", value=os.getenv("OPENAI_API_KEY", ""))
    runway_key = st.text_input("Runway API key", type="password", value=os.getenv("RUNWAYML_API_SECRET", ""))
    model = st.text_input("Reasoning model", value=DEFAULT_MODEL)
    scene_count = st.slider("Scenes", 6, 12, 8)
    st.markdown("**V2 rule:** no narration / no TTS.")

uploaded = st.file_uploader("Upload forecast PDF or XLSX", type=["pdf", "xlsx"])

if uploaded:
    data = uploaded.getvalue()
    try:
        source_text = extract_pdf(data) if uploaded.name.lower().endswith(".pdf") else extract_xlsx(data)
        source_text = clean_text(source_text)
    except Exception as e:
        st.error(f"Extraction failed: {e}")
        st.stop()

    st.success(f"Extracted {len(source_text.split()):,} words.")
    with st.expander("Source forecast"):
        st.text_area("Extracted text", source_text, height=280)

    if st.button("Run agentic interpretation", type="primary"):
        if not openai_key:
            st.error("Add an OpenAI API key.")
        else:
            with st.spinner("Analyst → Creative Director → Storyboard → Critic"):
                try:
                    client = OpenAI(api_key=openai_key)
                    analysis, creative, storyboard, critique = run_agents(
                        client, model, source_text, scene_count
                    )
                    st.session_state["analysis"] = analysis
                    st.session_state["creative"] = creative
                    st.session_state["storyboard"] = storyboard
                    st.session_state["critique"] = critique
                except Exception as e:
                    st.error(f"Agent pipeline failed: {e}")

if "storyboard" in st.session_state:
    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Forecast interpretation")
        st.json(st.session_state["analysis"])
    with c2:
        st.subheader("Creative direction")
        st.json(st.session_state["creative"])

    st.subheader("Critic")
    st.json(st.session_state["critique"])

    st.subheader("Human-in-the-loop storyboard")
    st.caption("Edit the AI's interpretation before any video is generated. Your edits become the approved prompts.")
    df = storyboard_df(st.session_state["storyboard"])
    edited = st.data_editor(
        df,
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        column_config={
            "scene": st.column_config.NumberColumn(disabled=True),
            "duration_sec": st.column_config.NumberColumn(min_value=2, max_value=10, step=1),
        },
        key="v2_storyboard_editor",
    )
    st.session_state["approved_scenes"] = edited.to_dict("records")

    st.download_button(
        "Download approved storyboard JSON",
        json.dumps(
            {
                "analysis": st.session_state["analysis"],
                "creative_direction": st.session_state["creative"],
                "scenes": st.session_state["approved_scenes"],
                "critique": st.session_state["critique"],
            },
            indent=2,
            ensure_ascii=False,
        ),
        file_name="approved_foresight_storyboard.json",
        mime="application/json",
    )

    st.subheader("Generate visual clips")
    st.warning("Video generation uses paid Runway credits. Test 1–2 scenes before generating the whole film.")

    selected = st.multiselect(
        "Scenes to generate",
        options=[int(s["scene"]) for s in st.session_state["approved_scenes"]],
        default=[int(st.session_state["approved_scenes"][0]["scene"])] if st.session_state["approved_scenes"] else [],
    )

    if st.button("Generate selected Runway clips"):
        if not runway_key:
            st.error("Add a Runway API key.")
        elif not selected:
            st.error("Select at least one scene.")
        else:
            os.environ["RUNWAYML_API_SECRET"] = runway_key
            runway = RunwayML()
            workdir = Path(tempfile.mkdtemp(prefix="foresight_v2_"))
            outputs = []

            progress = st.progress(0)
            chosen = [s for s in st.session_state["approved_scenes"] if int(s["scene"]) in selected]
            for i, scene in enumerate(chosen, 1):
                try:
                    with st.spinner(f"Generating scene {scene['scene']}..."):
                        url = generate_runway_clip(
                            runway,
                            scene["runway_prompt"],
                            scene.get("duration_sec", 5),
                        )
                        path = workdir / f"scene_{int(scene['scene']):02d}.mp4"
                        download(url, path)
                        outputs.append(path)
                        st.markdown(f"**Scene {scene['scene']} — {scene['forecast_anchor']}**")
                        st.video(path.read_bytes())
                except TaskFailedError as e:
                    st.error(f"Scene {scene['scene']} failed: {e}")
                except Exception as e:
                    st.error(f"Scene {scene['scene']} failed: {e}")
                progress.progress(i / len(chosen))

            if outputs:
                zip_path = workdir / "generated_foresight_clips.zip"
                with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
                    for p in outputs:
                        z.write(p, arcname=p.name)
                st.download_button(
                    "Download generated clips",
                    zip_path.read_bytes(),
                    file_name="generated_foresight_clips.zip",
                    mime="application/zip",
                )
else:
    if not uploaded:
        st.info("Upload a human, AI, or HITL forecast. The same pipeline can later be used for all three experimental conditions.")

st.divider()
st.caption("V2 generates abstract visual prompts and short generative-video clips. Final editing/music can be added after the visual language is validated.")
